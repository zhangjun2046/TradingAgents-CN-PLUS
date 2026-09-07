"""
报告导出工具 - 支持 Markdown、Word、PDF 格式

Word: 优先 pandoc，缺失或失败时降级 python-docx
PDF: WeasyPrint → pdfkit + wkhtmltopdf → fpdf2（纯 Python，适合 Windows）
"""

import logging
import os
import re
import tempfile
from io import BytesIO
from typing import Dict, Any, List, Optional

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# 可选依赖探测：任一失败不得阻断其它引擎
# ---------------------------------------------------------------------------

try:
    import markdown as markdown_lib
    MARKDOWN_AVAILABLE = True
except ImportError:
    markdown_lib = None
    MARKDOWN_AVAILABLE = False
    logger.warning("⚠️ markdown 未安装，PDF HTML 转换将受限")

try:
    import pypandoc
    try:
        pypandoc.get_pandoc_version()
        PANDOC_AVAILABLE = True
        logger.info("✅ Pandoc 可用")
    except OSError:
        PANDOC_AVAILABLE = False
        logger.warning("⚠️ Pandoc 不可用，Word 将尝试 python-docx 降级")
except ImportError:
    pypandoc = None
    PANDOC_AVAILABLE = False
    logger.warning("⚠️ pypandoc 未安装，Word 将尝试 python-docx 降级")

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    PYTHON_DOCX_AVAILABLE = True
    logger.info("✅ python-docx 可用")
except ImportError:
    Document = None
    qn = None
    Pt = None
    PYTHON_DOCX_AVAILABLE = False
    logger.warning("⚠️ python-docx 未安装，Word 降级导出不可用")

WEASYPRINT_AVAILABLE = False
WEASYPRINT_ERROR = None
try:
    from weasyprint import HTML as WeasyHTML
    WEASYPRINT_AVAILABLE = True
    logger.info("✅ WeasyPrint 可用（PDF 生成工具）")
except Exception as e:
    WeasyHTML = None
    WEASYPRINT_ERROR = str(e)
    logger.warning(f"⚠️ WeasyPrint 不可用: {e}")

PDFKIT_AVAILABLE = False
PDFKIT_ERROR = None
try:
    import pdfkit
    try:
        pdfkit.configuration()
        PDFKIT_AVAILABLE = True
        logger.info("✅ pdfkit + wkhtmltopdf 可用（PDF 生成工具）")
    except Exception as e:
        PDFKIT_ERROR = str(e)
        logger.warning("⚠️ wkhtmltopdf 未安装，pdfkit 不可用")
except ImportError:
    pdfkit = None
    logger.warning("⚠️ pdfkit 未安装")
except Exception as e:
    PDFKIT_ERROR = str(e)
    logger.warning(f"⚠️ pdfkit 检测失败: {e}")

FPDF_AVAILABLE = False
FPDF_ERROR = None
try:
    from fpdf import FPDF
    FPDF_AVAILABLE = True
    logger.info("✅ fpdf2 可用（纯 Python PDF 降级引擎）")
except Exception as e:
    FPDF = None
    FPDF_ERROR = str(e)
    logger.warning(f"⚠️ fpdf2 不可用: {e}")

EXPORT_AVAILABLE = MARKDOWN_AVAILABLE or PYTHON_DOCX_AVAILABLE or PANDOC_AVAILABLE


def _find_cjk_font() -> Optional[str]:
    """查找系统中文字体，供 fpdf2 嵌入 PDF。"""
    win_dir = os.environ.get("WINDIR", r"C:\Windows")
    candidates = [
        os.path.join(win_dir, "Fonts", "simhei.ttf"),
        os.path.join(win_dir, "Fonts", "simsun.ttc"),
        os.path.join(win_dir, "Fonts", "msyh.ttc"),
        os.path.join(win_dir, "Fonts", "msyhbd.ttc"),
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return None


CJK_FONT_PATH = _find_cjk_font()

# 实际入库的报告模块顺序与中文标题
MODULE_ORDER = [
    "company_overview",
    "financial_analysis",
    "technical_analysis",
    "market_analysis",
    "market_report",
    "fundamentals_report",
    "sentiment_report",
    "news_report",
    "risk_analysis",
    "valuation_analysis",
    "investment_plan",
    "investment_recommendation",
    "bull_researcher",
    "bear_researcher",
    "research_team_decision",
    "trader_investment_plan",
    "risky_analyst",
    "safe_analyst",
    "neutral_analyst",
    "risk_management_decision",
    "final_trade_decision",
]

MODULE_TITLES = {
    "company_overview": "公司概况",
    "financial_analysis": "财务分析",
    "technical_analysis": "技术分析",
    "market_analysis": "市场分析",
    "market_report": "市场技术分析",
    "fundamentals_report": "基本面分析",
    "sentiment_report": "情绪分析",
    "news_report": "新闻分析",
    "risk_analysis": "风险分析",
    "valuation_analysis": "估值分析",
    "investment_plan": "投资计划",
    "investment_recommendation": "投资建议",
    "bull_researcher": "多头研究员观点",
    "bear_researcher": "空头研究员观点",
    "research_team_decision": "研究团队决策",
    "trader_investment_plan": "交易员计划",
    "risky_analyst": "激进分析师观点",
    "safe_analyst": "保守分析师观点",
    "neutral_analyst": "中性分析师观点",
    "risk_management_decision": "风险管理决策",
    "final_trade_decision": "最终交易决策",
}


class ReportExporter:
    """报告导出器 - 支持 Markdown、Word、PDF 格式"""

    def __init__(self):
        self.export_available = EXPORT_AVAILABLE
        self.pandoc_available = PANDOC_AVAILABLE
        self.python_docx_available = PYTHON_DOCX_AVAILABLE
        self.weasyprint_available = WEASYPRINT_AVAILABLE
        self.pdfkit_available = PDFKIT_AVAILABLE
        self.fpdf_available = FPDF_AVAILABLE and bool(CJK_FONT_PATH)

        logger.info("📋 ReportExporter 初始化:")
        logger.info(f"  - export_available: {self.export_available}")
        logger.info(f"  - pandoc_available: {self.pandoc_available}")
        logger.info(f"  - python_docx_available: {self.python_docx_available}")
        logger.info(f"  - weasyprint_available: {self.weasyprint_available}")
        logger.info(f"  - pdfkit_available: {self.pdfkit_available}")
        logger.info(f"  - fpdf_available: {self.fpdf_available}")
        logger.info(f"  - cjk_font: {CJK_FONT_PATH}")

    @property
    def docx_available(self) -> bool:
        """Word 导出：pandoc 或 python-docx 任一可用即可"""
        return self.pandoc_available or self.python_docx_available

    @property
    def pdf_available(self) -> bool:
        """PDF 导出：WeasyPrint、pdfkit 或 fpdf2 任一可用即可"""
        return self.weasyprint_available or self.pdfkit_available or self.fpdf_available

    def generate_markdown_report(self, report_doc: Dict[str, Any]) -> str:
        """生成 Markdown 格式报告"""
        logger.info("📝 生成 Markdown 报告...")

        stock_symbol = report_doc.get("stock_symbol", "unknown")
        stock_name = report_doc.get("stock_name") or stock_symbol
        analysis_date = report_doc.get("analysis_date", "")
        analysts = report_doc.get("analysts", [])
        research_depth = report_doc.get("research_depth", 1)
        reports = report_doc.get("reports", {})
        summary = report_doc.get("summary", "")

        content_parts: List[str] = []

        title = f"{stock_name} 分析报告" if stock_name != stock_symbol else f"{stock_symbol} 股票分析报告"
        content_parts.append(f"# {title}")
        content_parts.append("")
        content_parts.append(f"**股票代码**: {stock_symbol}")
        content_parts.append(f"**分析日期**: {analysis_date}")
        if analysts:
            analyst_text = ", ".join(analysts) if isinstance(analysts, list) else str(analysts)
            content_parts.append(f"**分析师**: {analyst_text}")
        content_parts.append(f"**研究深度**: {research_depth}")
        content_parts.append("")
        content_parts.append("---")
        content_parts.append("")

        if summary:
            content_parts.append("## 执行摘要")
            content_parts.append("")
            content_parts.append(summary)
            content_parts.append("")
            content_parts.append("---")
            content_parts.append("")

        used_keys = set()
        for module_key in MODULE_ORDER:
            module_content = reports.get(module_key)
            if isinstance(module_content, str) and module_content.strip():
                title = MODULE_TITLES.get(module_key, module_key)
                content_parts.append(f"## {title}")
                content_parts.append("")
                content_parts.append(module_content.strip())
                content_parts.append("")
                content_parts.append("---")
                content_parts.append("")
                used_keys.add(module_key)

        for module_key, module_content in reports.items():
            if module_key in used_keys:
                continue
            if isinstance(module_content, str) and module_content.strip():
                title = MODULE_TITLES.get(module_key, module_key)
                content_parts.append(f"## {title}")
                content_parts.append("")
                content_parts.append(module_content.strip())
                content_parts.append("")
                content_parts.append("---")
                content_parts.append("")

        content_parts.append("")
        content_parts.append("---")
        content_parts.append("")
        content_parts.append("*本报告由 TradingAgents-CN 自动生成*")
        content_parts.append("")

        markdown_content = "\n".join(content_parts)
        logger.info(f"✅ Markdown 报告生成完成，长度: {len(markdown_content)} 字符")
        return markdown_content

    def _clean_markdown_for_pandoc(self, md_content: str) -> str:
        """清理 Markdown 内容，避免 pandoc 解析问题"""
        if md_content.strip().startswith("---"):
            md_content = "\n" + md_content

        md_content = re.sub(r"<[^>]*writing-mode[^>]*>", "", md_content, flags=re.IGNORECASE)
        md_content = re.sub(r"<[^>]*text-orientation[^>]*>", "", md_content, flags=re.IGNORECASE)
        md_content = re.sub(r'<div\s+style="[^"]*">', "<div>", md_content, flags=re.IGNORECASE)
        md_content = re.sub(r'<span\s+style="[^"]*">', "<span>", md_content, flags=re.IGNORECASE)
        md_content = re.sub(r"<style[^>]*>.*?</style>", "", md_content, flags=re.DOTALL | re.IGNORECASE)
        return md_content

    def generate_docx_report(self, report_doc: Dict[str, Any]) -> bytes:
        """生成 Word 文档：优先 pandoc，失败或不可用时降级 python-docx"""
        logger.info("📄 开始生成 Word 文档...")
        md_content = self.generate_markdown_report(report_doc)

        errors: List[str] = []
        if self.pandoc_available:
            try:
                return self._generate_docx_with_pandoc(md_content)
            except Exception as e:
                errors.append(f"pandoc: {e}")
                logger.warning(f"⚠️ pandoc 生成 Word 失败，尝试 python-docx 降级: {e}")

        if self.python_docx_available:
            try:
                return self._generate_docx_with_python_docx(md_content)
            except Exception as e:
                errors.append(f"python-docx: {e}")
                logger.error(f"❌ python-docx 生成 Word 失败: {e}", exc_info=True)

        detail = "；".join(errors) if errors else "未安装 pandoc 或 python-docx"
        raise Exception(
            f"无法生成 Word 文档。{detail}。"
            "请安装 python-docx，或安装 pandoc 后重试。"
        )

    def _generate_docx_with_pandoc(self, md_content: str) -> bytes:
        """使用 pypandoc 将 Markdown 转为 Word"""
        output_file = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
                output_file = tmp_file.name

            extra_args = [
                "--from=markdown-yaml_metadata_block",
                "--standalone",
                "--wrap=preserve",
                "--columns=120",
                "-M", "lang=zh-CN",
                "-M", "dir=ltr",
            ]
            cleaned_content = self._clean_markdown_for_pandoc(md_content)
            pypandoc.convert_text(
                cleaned_content,
                "docx",
                format="markdown",
                outputfile=output_file,
                extra_args=extra_args,
            )
            logger.info("✅ pypandoc 转换完成")
            self._fix_docx_text_direction(output_file)

            with open(output_file, "rb") as f:
                docx_content = f.read()
            logger.info(f"✅ Word 文档生成成功（pandoc），大小: {len(docx_content)} 字节")
            return docx_content
        finally:
            if output_file and os.path.exists(output_file):
                try:
                    os.unlink(output_file)
                except OSError:
                    pass

    def _fix_docx_text_direction(self, output_file: str) -> None:
        """后处理：移除 Word 文档中可能的竖排/双向文本设置"""
        if not PYTHON_DOCX_AVAILABLE:
            return
        try:
            doc = Document(output_file)
            for paragraph in doc.paragraphs:
                if paragraph._element.pPr is not None:
                    for child in list(paragraph._element.pPr):
                        if "textDirection" in child.tag or "bidi" in child.tag:
                            paragraph._element.pPr.remove(child)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            if paragraph._element.pPr is not None:
                                for child in list(paragraph._element.pPr):
                                    if "textDirection" in child.tag or "bidi" in child.tag:
                                        paragraph._element.pPr.remove(child)
            doc.save(output_file)
            logger.info("✅ Word 文档文本方向修复完成")
        except Exception as e:
            logger.warning(f"⚠️ Word 文档文本方向修复失败: {e}")

    def _set_run_font(self, run, font_name: str = "Microsoft YaHei") -> None:
        """设置中英文字体，避免中文显示异常"""
        run.font.name = font_name
        if run.font.size is None:
            run.font.size = Pt(11)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), font_name)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)

    def _add_inline_runs(self, paragraph, text: str) -> None:
        """将行内 Markdown（粗体/斜体/代码）写入 Word 段落"""
        pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\n]+?\*)")
        pos = 0
        for match in pattern.finditer(text):
            if match.start() > pos:
                run = paragraph.add_run(text[pos:match.start()])
                self._set_run_font(run)
            token = match.group(0)
            if token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
                self._set_run_font(run)
            elif token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                self._set_run_font(run, "Consolas")
            elif token.startswith("*") and token.endswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
                self._set_run_font(run)
            pos = match.end()
        if pos < len(text):
            run = paragraph.add_run(text[pos:])
            self._set_run_font(run)

    def _is_table_separator(self, line: str) -> bool:
        stripped = line.strip()
        if "|" not in stripped:
            return False
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        return all(re.match(r"^:?-{3,}:?$", c) for c in cells if c)

    def _split_table_row(self, line: str) -> List[str]:
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [cell.strip() for cell in line.split("|")]

    def _generate_docx_with_python_docx(self, md_content: str) -> bytes:
        """不依赖 pandoc，用 python-docx 将 Markdown 转为 Word"""
        logger.info("🔧 使用 python-docx 生成 Word...")
        doc = Document()

        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        lines = md_content.split("\n")
        i = 0
        in_code_block = False
        code_lines: List[str] = []

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            if stripped.startswith("```"):
                if in_code_block:
                    paragraph = doc.add_paragraph("\n".join(code_lines))
                    if paragraph.runs:
                        self._set_run_font(paragraph.runs[0], "Consolas")
                    else:
                        run = paragraph.add_run("\n".join(code_lines))
                        self._set_run_font(run, "Consolas")
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                i += 1
                continue

            if in_code_block:
                code_lines.append(line)
                i += 1
                continue

            if not stripped:
                i += 1
                continue

            if stripped in ("---", "***", "___"):
                doc.add_paragraph("─" * 20)
                i += 1
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if heading_match:
                level = min(len(heading_match.group(1)), 4)
                heading = doc.add_heading(heading_match.group(2).strip(), level=level)
                for run in heading.runs:
                    self._set_run_font(run)
                i += 1
                continue

            if stripped.startswith("|") and i + 1 < len(lines) and self._is_table_separator(lines[i + 1]):
                table_rows = [self._split_table_row(stripped)]
                i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    if not self._is_table_separator(lines[i]):
                        table_rows.append(self._split_table_row(lines[i]))
                    i += 1
                if table_rows:
                    col_count = max(len(r) for r in table_rows)
                    table = doc.add_table(rows=len(table_rows), cols=col_count)
                    table.style = "Table Grid"
                    for r_idx, row_cells in enumerate(table_rows):
                        for c_idx in range(col_count):
                            cell_text = row_cells[c_idx] if c_idx < len(row_cells) else ""
                            cell = table.cell(r_idx, c_idx)
                            cell.text = ""
                            paragraph = cell.paragraphs[0]
                            self._add_inline_runs(paragraph, cell_text)
                            if r_idx == 0:
                                for run in paragraph.runs:
                                    run.bold = True
                continue

            list_match = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
            if list_match:
                content = list_match.group(3)
                is_numbered = list_match.group(2)[-1] == "."
                style_name = "List Number" if is_numbered else "List Bullet"
                try:
                    paragraph = doc.add_paragraph(style=style_name)
                except (KeyError, ValueError):
                    paragraph = doc.add_paragraph()
                self._add_inline_runs(paragraph, content)
                i += 1
                continue

            paragraph = doc.add_paragraph()
            self._add_inline_runs(paragraph, stripped)
            i += 1

        buffer = BytesIO()
        doc.save(buffer)
        docx_content = buffer.getvalue()
        logger.info(f"✅ Word 文档生成成功（python-docx），大小: {len(docx_content)} 字节")
        return docx_content

    def _markdown_body_html(self, md_content: str) -> str:
        """将 Markdown 转为 HTML 片段（不含完整文档壳）"""
        if MARKDOWN_AVAILABLE:
            extensions = [
                "markdown.extensions.tables",
                "markdown.extensions.fenced_code",
                "markdown.extensions.nl2br",
            ]
            return markdown_lib.markdown(md_content, extensions=extensions)
        escaped = (
            md_content.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        return f"<pre>{escaped}</pre>"

    def _markdown_to_html(self, md_content: str) -> str:
        """将 Markdown 转换为带中文样式的 HTML"""
        html_content = self._markdown_body_html(md_content)

        return f"""
<!DOCTYPE html>
<html lang="zh-CN" dir="ltr">
<head>
    <meta charset="UTF-8">
    <title>分析报告</title>
    <style>
        html {{ direction: ltr; }}
        body {{
            font-family: "Noto Sans CJK SC", "Source Han Sans SC", "Microsoft YaHei",
                         "SimHei", "WenQuanYi Micro Hei", "Arial", sans-serif;
            line-height: 1.8;
            color: #333;
            margin: 20mm;
            padding: 0;
            background: white;
            direction: ltr;
        }}
        h1, h2, h3, h4, h5, h6 {{
            color: #2c3e50;
            margin-top: 1.5em;
            margin-bottom: 0.8em;
            font-weight: 600;
            page-break-after: avoid;
            direction: ltr;
        }}
        h1 {{
            font-size: 2em;
            border-bottom: 3px solid #3498db;
            padding-bottom: 0.3em;
            page-break-before: always;
        }}
        h1:first-child {{ page-break-before: avoid; }}
        h2 {{
            font-size: 1.6em;
            border-bottom: 2px solid #bdc3c7;
            padding-bottom: 0.25em;
        }}
        h3 {{ font-size: 1.3em; color: #34495e; }}
        p {{ margin: 0.8em 0; text-align: left; direction: ltr; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 1.5em 0;
            font-size: 0.9em;
            direction: ltr;
        }}
        thead {{ display: table-header-group; }}
        tbody {{ display: table-row-group; }}
        tr {{ page-break-inside: avoid; }}
        th, td {{
            border: 1px solid #ddd;
            padding: 10px 12px;
            text-align: left;
            direction: ltr;
        }}
        th {{ background-color: #3498db; color: white; font-weight: bold; }}
        tbody tr:nth-child(even) {{ background-color: #f8f9fa; }}
        code {{
            background-color: #f4f4f4;
            padding: 2px 6px;
            border-radius: 3px;
            font-family: "Consolas", "Monaco", "Courier New", monospace;
            font-size: 0.9em;
        }}
        pre {{
            background-color: #f4f4f4;
            padding: 15px;
            border-radius: 5px;
            border-left: 4px solid #3498db;
            page-break-inside: avoid;
            white-space: pre-wrap;
            word-wrap: break-word;
        }}
        pre code {{ background-color: transparent; padding: 0; }}
        ul, ol {{ margin: 0.8em 0; padding-left: 2em; direction: ltr; }}
        li {{ margin: 0.4em 0; }}
        strong, b {{ font-weight: 700; color: #2c3e50; }}
        hr {{ border: none; border-top: 2px solid #ecf0f1; margin: 2em 0; }}
        blockquote {{
            margin: 1em 0;
            padding: 0.5em 1em;
            border-left: 4px solid #3498db;
            background-color: #f8f9fa;
            page-break-inside: avoid;
        }}
        img {{ max-width: 100%; height: auto; page-break-inside: avoid; }}
        @page {{
            size: A4;
            margin: 20mm;
        }}
        p, li {{ orphans: 3; widows: 3; }}
    </style>
</head>
<body>
{html_content}
</body>
</html>
"""

    def _generate_pdf_with_weasyprint(self, html_content: str) -> bytes:
        """使用 WeasyPrint 从 HTML 生成 PDF"""
        logger.info("🔧 使用 WeasyPrint 生成 PDF...")
        pdf_bytes = WeasyHTML(string=html_content).write_pdf()
        logger.info(f"✅ WeasyPrint PDF 生成成功，大小: {len(pdf_bytes)} 字节")
        return pdf_bytes

    def _generate_pdf_with_pdfkit(self, html_content: str) -> bytes:
        """使用 pdfkit + wkhtmltopdf 生成 PDF"""
        logger.info("🔧 使用 pdfkit + wkhtmltopdf 生成 PDF...")
        options = {
            "encoding": "UTF-8",
            "enable-local-file-access": None,
            "page-size": "A4",
            "margin-top": "20mm",
            "margin-right": "20mm",
            "margin-bottom": "20mm",
            "margin-left": "20mm",
        }
        pdf_bytes = pdfkit.from_string(html_content, False, options=options)
        logger.info(f"✅ pdfkit PDF 生成成功，大小: {len(pdf_bytes)} 字节")
        return pdf_bytes

    def _fpdf_reset_x(self, pdf) -> None:
        pdf.set_x(pdf.l_margin)

    def _fpdf_text(self, pdf, text: str, height: float = 7) -> None:
        """写一段文本，避免表格/横线后 x 坐标未回页边导致无空间绘图"""
        self._fpdf_reset_x(pdf)
        content = text if text else " "
        try:
            pdf.multi_cell(0, height, content)
        except Exception:
            pdf.multi_cell(pdf.epw, height, content)
        self._fpdf_reset_x(pdf)

    def _plain_text(self, text: str) -> str:
        """去掉 Markdown/HTML 标记，供 fpdf 纯文本绘制"""
        if not text:
            return ""
        text = re.sub(r"```.*?```", " ", text, flags=re.DOTALL)
        text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"</p>\s*<p[^>]*>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`(.+?)`", r"\1", text)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        text = text.replace("&nbsp;", " ").replace("&amp;", "&")
        text = text.replace("&lt;", "<").replace("&gt;", ">")
        text = text.replace("•", "·").replace("&#8226;", "·")
        return re.sub(r"[ \t]+", " ", text).strip()

    def _html_table_to_rows(self, table_html: str) -> List[List[str]]:
        """把 HTML 表格拍成二维文本，去掉嵌套 td/table"""
        rows: List[List[str]] = []
        for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", table_html, flags=re.IGNORECASE | re.DOTALL):
            # 只取当前行的直接单元格，避免把内层表格的 td 算进来
            cells = re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, flags=re.IGNORECASE | re.DOTALL)
            cleaned = []
            for cell in cells:
                cell = re.sub(
                    r"<table.*?</table>",
                    lambda m: " ".join(self._plain_text(x) for x in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", m.group(0), flags=re.I | re.S)),
                    cell,
                    flags=re.IGNORECASE | re.DOTALL,
                )
                cleaned.append(self._plain_text(cell))
            if cleaned:
                rows.append(cleaned)
        return rows

    def _render_fpdf_table(self, pdf, rows: List[List[str]], font_name: str) -> None:
        """用 fpdf2 原生表格绘制，避免 write_html 不支持嵌套 td"""
        if not rows:
            return
        col_count = max(len(r) for r in rows)
        padded = [list(r) + [""] * (col_count - len(r)) for r in rows]
        pdf.set_font(font_name, size=9)
        try:
            with pdf.table(width=pdf.epw, line_height=6, text_align="LEFT") as table:
                for data_row in padded:
                    row = table.row()
                    for datum in data_row:
                        row.cell(self._plain_text(str(datum)) or " ")
        except Exception as e:
            logger.warning(f"⚠️ fpdf 原生表格失败，降级为纯文本: {e}")
            for data_row in padded:
                self._fpdf_text(pdf, " | ".join(self._plain_text(str(c)) for c in data_row), height=6)
        pdf.set_font(font_name, size=11)
        self._fpdf_reset_x(pdf)
        pdf.ln(2)

    def _generate_pdf_with_fpdf(self, md_content: str) -> bytes:
        """使用 fpdf2 + 系统中文字体生成 PDF（不走 HTML，兼容复杂表格）"""
        if FPDF is None:
            raise Exception("fpdf2 未安装")
        if not CJK_FONT_PATH:
            raise Exception("未找到中文字体，无法使用 fpdf2 生成 PDF")

        logger.info(f"🔧 使用 fpdf2 生成 PDF，字体: {CJK_FONT_PATH}")
        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        font_name = "CJK"
        pdf.add_font(font_name, "", CJK_FONT_PATH)
        pdf.add_font(font_name, "B", CJK_FONT_PATH)
        pdf.add_font(font_name, "I", CJK_FONT_PATH)
        pdf.set_font(font_name, size=11)

        lines = md_content.split("\n")
        i = 0
        in_code_block = False
        code_lines: List[str] = []
        heading_sizes = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            if stripped.startswith("```"):
                if in_code_block:
                    self._fpdf_text(pdf, "\n".join(code_lines) or " ", height=5)
                    pdf.ln(2)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                i += 1
                continue

            if in_code_block:
                code_lines.append(line)
                i += 1
                continue

            if not stripped:
                i += 1
                continue

            if stripped.lower().startswith("<table"):
                html_buf = [line]
                if "</table>" not in line.lower():
                    i += 1
                    while i < len(lines) and "</table>" not in lines[i].lower():
                        html_buf.append(lines[i])
                        i += 1
                    if i < len(lines):
                        html_buf.append(lines[i])
                        i += 1
                else:
                    i += 1
                self._render_fpdf_table(pdf, self._html_table_to_rows("\n".join(html_buf)), font_name)
                continue

            if stripped in ("---", "***", "___"):
                self._fpdf_reset_x(pdf)
                pdf.ln(1)
                y = pdf.get_y()
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(3)
                self._fpdf_reset_x(pdf)
                i += 1
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if heading_match:
                level = min(len(heading_match.group(1)), 6)
                pdf.set_font(font_name, "B", heading_sizes.get(level, 12))
                self._fpdf_text(pdf, self._plain_text(heading_match.group(2)), height=9)
                pdf.set_font(font_name, size=11)
                pdf.ln(1)
                i += 1
                continue

            if stripped.startswith("|") and i + 1 < len(lines) and self._is_table_separator(lines[i + 1]):
                table_rows = [self._split_table_row(stripped)]
                i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    if not self._is_table_separator(lines[i]):
                        table_rows.append(self._split_table_row(lines[i]))
                    i += 1
                self._render_fpdf_table(pdf, table_rows, font_name)
                continue

            list_match = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
            if list_match:
                self._fpdf_text(pdf, "· " + self._plain_text(list_match.group(3)))
                i += 1
                continue

            self._fpdf_text(pdf, self._plain_text(stripped))
            i += 1

        buffer = BytesIO()
        pdf.output(buffer)
        pdf_bytes = buffer.getvalue()
        logger.info(f"✅ fpdf2 PDF 生成成功，大小: {len(pdf_bytes)} 字节")
        return pdf_bytes

    def generate_pdf_report(self, report_doc: Dict[str, Any]) -> bytes:
        """生成 PDF：WeasyPrint → pdfkit → fpdf2"""
        logger.info("📊 开始生成 PDF 文档...")

        if not self.pdf_available:
            hints = []
            if WEASYPRINT_ERROR:
                hints.append(f"WeasyPrint: {WEASYPRINT_ERROR}")
            if PDFKIT_ERROR:
                hints.append(f"pdfkit: {PDFKIT_ERROR}")
            if FPDF_ERROR:
                hints.append(f"fpdf2: {FPDF_ERROR}")
            if FPDF_AVAILABLE and not CJK_FONT_PATH:
                hints.append("未找到中文字体（Windows 需有 simhei.ttf 或 msyh.ttc）")
            extra = ("\n" + "\n".join(hints)) if hints else ""
            raise Exception(
                "PDF 导出功能不可用。请安装 fpdf2（pip install fpdf2），"
                "或安装 WeasyPrint / pdfkit + wkhtmltopdf。" + extra
            )

        md_content = self.generate_markdown_report(report_doc)
        html_content = self._markdown_to_html(md_content)
        errors: List[str] = []

        if self.weasyprint_available:
            try:
                return self._generate_pdf_with_weasyprint(html_content)
            except Exception as e:
                errors.append(f"WeasyPrint: {e}")
                logger.warning(f"⚠️ WeasyPrint 生成 PDF 失败，尝试下一引擎: {e}")

        if self.pdfkit_available:
            try:
                return self._generate_pdf_with_pdfkit(html_content)
            except Exception as e:
                errors.append(f"pdfkit: {e}")
                logger.warning(f"⚠️ pdfkit 生成 PDF 失败，尝试 fpdf2: {e}")

        if self.fpdf_available:
            try:
                return self._generate_pdf_with_fpdf(md_content)
            except Exception as e:
                errors.append(f"fpdf2: {e}")
                logger.error(f"❌ fpdf2 生成 PDF 失败: {e}")

        detail = "；".join(errors) if errors else "没有可用的 PDF 引擎"
        raise Exception(
            f"PDF 生成失败: {detail}。"
            "请安装 fpdf2，或安装 WeasyPrint / wkhtmltopdf 后重试。"
        )


report_exporter = ReportExporter()
